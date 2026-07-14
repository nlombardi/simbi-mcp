"""Tests for CSV/Excel structure profiling (analyze_data_source's core)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl
import polars as pl
import pytest

from simbi_mcp.semantic.data_profile import (
    _column_hints,
    _map_dtype_to_tmdl,
    _table_hints,
    profile_csv,
    profile_dataframe,
    profile_excel,
    profile_file,
)
from simbi_mcp.types import ColumnProfile


def _find(columns: list[ColumnProfile], name: str) -> ColumnProfile:
    return next(c for c in columns if c.name == name)


@pytest.mark.parametrize(
    "dtype,expected",
    [
        (pl.Int64, "int64"),
        (pl.Int32, "int64"),
        (pl.UInt8, "int64"),
        (pl.Float64, "double"),
        (pl.Float32, "double"),
        (pl.Utf8, "string"),
        (pl.Boolean, "boolean"),
        (pl.Date, "dateTime"),
        (pl.Datetime, "dateTime"),
    ],
)
def test_map_dtype_to_tmdl(dtype: pl.DataType, expected: str) -> None:
    assert _map_dtype_to_tmdl(dtype) == expected


def test_map_dtype_to_tmdl_falls_back_to_string_for_nested_types() -> None:
    assert _map_dtype_to_tmdl(pl.List(pl.Int64)) == "string"


def test_profile_dataframe_row_count() -> None:
    df = pl.DataFrame({"OrderID": [1, 2, 3, 4]})
    profile = profile_dataframe(df, "sales")
    assert profile.table_name == "sales"
    assert profile.row_count == 4


def test_profile_dataframe_column_stats() -> None:
    df = pl.DataFrame(
        {
            "OrderID": [1, 2, 3, 4],
            "Region": ["North", "South", "North", None],
            "Revenue": [10.5, 20.0, None, 5.25],
        }
    )
    profile = profile_dataframe(df, "sales")

    order_id = _find(profile.columns, "OrderID")
    assert order_id.tmdl_type == "int64"
    assert order_id.null_count == 0
    assert order_id.distinct_count == 4
    assert order_id.min == "1"
    assert order_id.max == "4"

    region = _find(profile.columns, "Region")
    assert region.tmdl_type == "string"
    assert region.null_count == 1
    assert region.null_pct == 25.0
    assert region.distinct_count == 2  # non-null distinct: North, South
    assert region.min is None  # min/max only computed for numeric/temporal
    assert region.max is None
    assert set(region.sample_values) <= {"North", "South"}

    revenue = _find(profile.columns, "Revenue")
    assert revenue.tmdl_type == "double"
    assert revenue.null_count == 1
    assert revenue.min == "5.25"
    assert revenue.max == "20.0"


def test_primary_key_hint_fires_on_unique_id_named_column() -> None:
    hints = _column_hints("OrderID", pl.Int64, distinct_count=50, row_count=50)
    assert "likely primary key" in hints


def test_primary_key_hint_does_not_fire_without_id_like_name() -> None:
    hints = _column_hints("OrderDate", pl.Date, distinct_count=50, row_count=50)
    assert "likely primary key" not in hints


def test_primary_key_hint_does_not_fire_if_not_fully_unique() -> None:
    hints = _column_hints("CustomerID", pl.Int64, distinct_count=40, row_count=50)
    assert "likely primary key" not in hints


def test_date_hint_fires_for_date_dtype() -> None:
    hints = _column_hints("OrderDate", pl.Date, distinct_count=50, row_count=50)
    assert "time-intelligence candidate" in hints


def test_date_hint_fires_for_datetime_dtype() -> None:
    hints = _column_hints("CreatedAt", pl.Datetime("us"), distinct_count=50, row_count=50)
    assert "time-intelligence candidate" in hints


def test_dimension_hint_fires_for_low_cardinality_text() -> None:
    hints = _column_hints("Region", pl.Utf8, distinct_count=4, row_count=50)
    assert "dimension/slicer candidate" in hints


def test_high_cardinality_hint_fires_for_mostly_unique_text() -> None:
    hints = _column_hints("Description", pl.Utf8, distinct_count=95, row_count=100)
    assert "high cardinality — avoid as a slicer" in hints
    assert "dimension/slicer candidate" not in hints


def test_no_hint_fires_for_plain_numeric_measure() -> None:
    hints = _column_hints("Revenue", pl.Float64, distinct_count=48, row_count=50)
    assert hints == []


def test_wide_format_hint_fires_for_year_columns() -> None:
    hints = _table_hints(["Indicator", "2023", "2024", "2025"])
    assert len(hints) == 1
    assert "WIDE format" in hints[0]
    assert "2023" in hints[0] and "2024" in hints[0] and "2025" in hints[0]


def test_wide_format_hint_fires_for_quarter_columns() -> None:
    hints = _table_hints(["Region", "Q1", "Q2", "Q3", "Q4"])
    assert len(hints) == 1
    assert "WIDE format" in hints[0]


def test_wide_format_hint_does_not_fire_for_normal_columns() -> None:
    assert _table_hints(["OrderID", "OrderDate", "Region", "Revenue"]) == []


def test_wide_format_hint_requires_at_least_three_period_columns() -> None:
    assert _table_hints(["Indicator", "2023", "2024"]) == []


@pytest.fixture
def fixtures_datasets() -> Path:
    return Path(__file__).parent.parent / "fixtures" / "datasets"


def test_profile_csv_row_count_and_table_name(fixtures_datasets: Path) -> None:
    profile = profile_csv(fixtures_datasets / "sales_small.csv")
    assert profile.table_name == "sales_small"
    assert profile.row_count == 50


def test_profile_csv_primary_key_hint(fixtures_datasets: Path) -> None:
    profile = profile_csv(fixtures_datasets / "sales_small.csv")
    order_id = _find(profile.columns, "OrderID")
    assert order_id.tmdl_type == "int64"
    assert "likely primary key" in order_id.hints


def test_profile_csv_date_hint(fixtures_datasets: Path) -> None:
    # try_parse_dates=True is required for this to fire — polars otherwise
    # leaves ISO date strings as plain text, silently defeating the hint.
    profile = profile_csv(fixtures_datasets / "sales_small.csv")
    order_date = _find(profile.columns, "OrderDate")
    assert order_date.tmdl_type == "dateTime"
    assert "time-intelligence candidate" in order_date.hints


def test_profile_csv_dimension_hint(fixtures_datasets: Path) -> None:
    profile = profile_csv(fixtures_datasets / "sales_small.csv")
    region = _find(profile.columns, "Region")
    assert region.distinct_count == 4
    assert "dimension/slicer candidate" in region.hints


def test_profile_csv_no_wide_format_false_positive(fixtures_datasets: Path) -> None:
    profile = profile_csv(fixtures_datasets / "sales_small.csv")
    assert profile.hints == []


def test_profile_csv_missing_file_raises(fixtures_datasets: Path) -> None:
    with pytest.raises(ValueError, match="Could not read CSV"):
        profile_csv(fixtures_datasets / "does_not_exist.csv")


@pytest.fixture
def multi_sheet_xlsx(tmp_path: Path) -> Path:
    path = tmp_path / "multi.xlsx"
    workbook = openpyxl.Workbook()

    orders = workbook.active
    orders.title = "Orders"
    orders.append(["OrderID", "OrderDate", "Region"])
    orders.append([1, date(2025, 1, 1), "North"])
    orders.append([2, date(2025, 1, 2), "South"])
    orders.append([3, date(2025, 1, 3), "East"])

    notes = workbook.create_sheet("Notes")
    notes.append(["Comment"])
    notes.append(["first comment"])
    notes.append(["second comment"])

    workbook.save(path)
    return path


def test_profile_excel_reads_every_sheet_by_default(multi_sheet_xlsx: Path) -> None:
    tables = profile_excel(multi_sheet_xlsx)
    assert [t.table_name for t in tables] == ["Orders", "Notes"]
    assert tables[0].row_count == 3
    assert [c.name for c in tables[0].columns] == ["OrderID", "OrderDate", "Region"]


def test_profile_excel_filters_to_requested_sheet(multi_sheet_xlsx: Path) -> None:
    tables = profile_excel(multi_sheet_xlsx, sheet="Notes")
    assert len(tables) == 1
    assert tables[0].table_name == "Notes"
    assert tables[0].row_count == 2


def test_profile_excel_unknown_sheet_raises(multi_sheet_xlsx: Path) -> None:
    with pytest.raises(ValueError, match=r"Orders.*Notes|Notes.*Orders"):
        profile_excel(multi_sheet_xlsx, sheet="DoesNotExist")


def test_profile_excel_date_column_gets_hint(multi_sheet_xlsx: Path) -> None:
    tables = profile_excel(multi_sheet_xlsx, sheet="Orders")
    order_date = _find(tables[0].columns, "OrderDate")
    assert order_date.tmdl_type == "dateTime"
    assert "time-intelligence candidate" in order_date.hints


def test_profile_excel_corrupted_file_raises(tmp_path: Path) -> None:
    corrupted = tmp_path / "corrupted.xlsx"
    corrupted.write_bytes(b"this is not a real xlsx file")
    with pytest.raises(ValueError, match="Could not read Excel"):
        profile_excel(corrupted)


def test_profile_file_csv(fixtures_datasets: Path) -> None:
    csv_path = fixtures_datasets / "sales_small.csv"
    profile = profile_file(str(csv_path))
    assert profile.source_path == str(csv_path)
    assert len(profile.tables) == 1
    assert profile.tables[0].table_name == "sales_small"


def test_profile_file_xlsx(multi_sheet_xlsx: Path) -> None:
    profile = profile_file(str(multi_sheet_xlsx))
    assert [t.table_name for t in profile.tables] == ["Orders", "Notes"]


def test_profile_file_xlsx_with_sheet_filter(multi_sheet_xlsx: Path) -> None:
    profile = profile_file(str(multi_sheet_xlsx), sheet="Orders")
    assert len(profile.tables) == 1
    assert profile.tables[0].table_name == "Orders"


def test_profile_file_missing_file_raises() -> None:
    with pytest.raises(ValueError, match="File not found"):
        profile_file("does_not_exist_anywhere.csv")


def test_profile_file_unsupported_extension_raises(tmp_path: Path) -> None:
    bad_file = tmp_path / "data.txt"
    bad_file.write_text("hello", encoding="utf-8")
    with pytest.raises(ValueError, match=r"\.csv.*\.xlsx|\.xlsx.*\.csv"):
        profile_file(str(bad_file))


def test_profile_file_xls_gets_clear_unsupported_message(tmp_path: Path) -> None:
    bad_file = tmp_path / "legacy.xls"
    bad_file.write_text("not a real xls", encoding="utf-8")
    with pytest.raises(ValueError, match=r"(?i)not supported"):
        profile_file(str(bad_file))
